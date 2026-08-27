from tests.common import *


class TestAPIClient(unittest.TestCase):
    def test_default_config_free_api(self):
        client = IPAPIClient()
        self.assertIn('ip-api.com/json/', client.API_URL)
        self.assertNotIn('pro.', client.API_URL)
        self.assertEqual(client.batch_size, 45)
        self.assertEqual(client.period, 60)
        self.assertIsNone(client.api_key)

    def test_with_api_key(self):
        client = IPAPIClient(api_key='test_key_123')
        self.assertIn('pro.ip-api.com', client.API_URL)
        self.assertEqual(client.batch_size, 500)
        self.assertEqual(client.period, 0)
        self.assertEqual(client.api_key, 'test_key_123')
        self.assertIn('https', client.API_URL)

    def test_cache_init(self):
        client = IPAPIClient(cache_file=None)
        self.assertEqual(len(client.cache), 0)

    def test_make_error_result(self):
        client = IPAPIClient()
        result = client._make_error_result('test error')
        self.assertIn('Erro: test error', result['Ip_Dono'])
        self.assertIn('Ip_Pais', result)
        self.assertIn('Ip_Pais_Codigo', result)

    def test_cache_ttl_default(self):
        self.assertEqual(IPAPIClient().cache_ttl, 30 * 24 * 60 * 60)

    def test_cache_ttl_custom(self):
        self.assertEqual(IPAPIClient(cache_ttl=3600).cache_ttl, 3600)


class TestAuditLogger(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()
        self.audit_file = os.path.join(self.tmp_dir, 'test_audit.jsonl')

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_log_event(self):
        event = log_audit_event(action='test_action', details={'key': 'value'}, user='test_user')
        self.assertEqual(event['action'], 'test_action')
        self.assertEqual(event['user'], 'test_user')
        self.assertIn('timestamp', event)
        self.assertIn('event_hash', event)

    def test_integrity_receipt(self):
        receipt = generate_integrity_receipt(input_file=__file__, output_file=__file__, alvo='test_target', record_count=1)
        self.assertIn('receipt_hash', receipt)
        self.assertIn('record_count', receipt)
        self.assertEqual(receipt['record_count'], 1)


class TestReportGenerator(unittest.TestCase):
    def _make_sample_df(self):
        return pd.DataFrame({
            'Ip': ['8.8.8.8', '1.1.1.1', '8.8.8.8'],
            'Data': ['2025-01-01 10:00:00', '2025-01-01 11:00:00', '2025-01-02 12:00:00'],
            'Ip_Pais': ['US', 'US', 'US'],
            'Ip_Regiao': ['CA', 'CA', 'CA'],
            'Ip_Cidade': ['LA', 'LA', 'LA'],
            'Ip_Dono': ['Google', 'Cloudflare', 'Google'],
            'Ip_Proxy': [False, False, False],
            'Ip_Hospedagem': [False, False, False],
            'Ip_Movel': [False, False, False],
            'Periodo': ['Diurno', 'Diurno', 'Diurno'],
        })

    def test_generate_html_report(self):
        html_bytes = generate_html_report(self._make_sample_df(), 'test_target')
        self.assertIsInstance(html_bytes, (bytes, bytearray))
        self.assertGreater(len(html_bytes), 500)
        html_text = html_bytes.decode('utf-8')
        self.assertIn('<!DOCTYPE html>', html_text)
        self.assertIn('test_target', html_text)

    def test_html_report_with_config(self):
        config = {
            'title': 'Test Report',
            'organization': 'Test Org',
            'case_number': 'CASE-001',
            'analyst': 'Agent Smith',
            'classification': 'CONFIDENCIAL',
        }
        html_bytes = generate_html_report(self._make_sample_df(), 'target', config=config)
        html_text = html_bytes.decode('utf-8')
        self.assertIn('CONFIDENCIAL', html_text)
        self.assertIn('Agent Smith', html_text)
        self.assertIn('CASE-001', html_text)

    def test_html_report_with_analyses(self):
        analyses = {
            'risk_scores': pd.DataFrame({'IP': ['8.8.8.8'], 'Score': [75], 'Ip_Dono': ['Google'], 'Ip_Pais': ['US']}),
            'impossible_jumps': pd.DataFrame(),
            'vpn_heuristics': {'score': 0, 'indicators': {}, 'suspicious_ips': []},
        }
        html_bytes = generate_html_report(self._make_sample_df(), 'target', analyses=analyses)
        self.assertGreater(len(html_bytes), 500)

    def test_html_report_with_audit_hash(self):
        html_bytes = generate_html_report(self._make_sample_df(), 'target', audit_hash='abc123def456' * 5)
        html_text = html_bytes.decode('utf-8')
        self.assertIn('abc123def456', html_text)


class TestCacheCompression(unittest.TestCase):
    def test_cache_metrics(self):
        client = IPAPIClient(cache_file=None)
        self.assertEqual(client.cache_hits, 0)
        self.assertEqual(client.cache_misses, 0)

    def test_get_status(self):
        status = IPAPIClient(cache_file=None).get_status()
        self.assertIn('cache_hit_rate', status)
        self.assertIn('cache_size', status)


class TestSmokeFlows(unittest.TestCase):
    def _make_enriched_df(self):
        df = extrair_ips_do_formato_simples('8.8.8.8\n1.1.1.1', alvo='smoke')
        df['Data'] = ['2025-01-01 10:00:00', '2025-01-01 11:00:00']
        df['Ip_Pais'] = ['BR', 'BR']
        df['Ip_Regiao'] = ['SP', 'SP']
        df['Ip_Cidade'] = ['Sao Paulo', 'Sao Paulo']
        df['Ip_Dono'] = ['Google', 'Cloudflare']
        df['Ip_Proxy'] = [False, False]
        df['Ip_Hospedagem'] = [False, False]
        df['Ip_Movel'] = [False, False]
        df['Periodo'] = ['☀️ Diurno', '☀️ Diurno']
        return df

    def test_simple_input_report_and_xlsx_smoke(self):
        from file_handler import export_xlsx_colored

        df = self._make_enriched_df()
        html_bytes = generate_html_report(df, 'smoke')
        self.assertGreater(len(html_bytes), 500)
        self.assertIn(b'<!DOCTYPE html>', html_bytes)

        output = io.BytesIO()
        export_xlsx_colored(df, output)
        self.assertGreater(len(output.getvalue()), 0)

    def test_interception_enrichment_smoke(self):
        from interception_parser import processar_resultados_interceptacao

        df = pd.DataFrame({
            'Sender IP': ['8.8.8.8'],
            'Ip_Dono': [''],
            'Ip_AS': [''],
            'Ip_Regiao': [''],
            'Ip_Cidade': [''],
            'Ip_Pais': [''],
            'Ip_Pais_Codigo': [''],
            'Ip_Movel': [False],
            'Ip_Proxy': [False],
            'Ip_Hospedagem': [False],
            'Ip_Lat': [None],
            'Ip_Lon': [None],
        })
        resultados_api = {
            '8.8.8.8': {
                'status': 'success',
                'Ip_Dono': 'Google',
                'Ip_AS': 'AS15169',
                'Ip_Regiao': 'California',
                'Ip_Cidade': 'Mountain View',
                'Ip_Pais': 'United States',
                'Ip_Pais_Codigo': 'US',
                'Ip_Movel': False,
                'Ip_Proxy': False,
                'Ip_Hospedagem': False,
                'Ip_Lat': 37.4056,
                'Ip_Lon': -122.0775,
            }
        }
        result = processar_resultados_interceptacao(df, resultados_api)
        self.assertEqual(result.loc[0, 'Ip_Dono'], 'Google')
        self.assertEqual(result.loc[0, 'Ip_Pais_Codigo'], 'US')


class TestMapaSemChaveDeApi(unittest.TestCase):
    """Os tiles da CARTO passaram a exigir chave: a requisição volta 200 com a
    imagem carimbada "API KEY REQUIRED" por cima do mapa inteiro. Nada falha em
    voz alta — o laudo simplesmente sai com a figura inutilizada."""

    def test_relatorio_nao_usa_tiles_da_carto(self):
        from html_report_generator import generate_html_report

        df = pd.DataFrame({
            'Alvo': ['t'] * 2,
            'Ip': ['198.51.100.1', '198.51.100.2'],
            'Data': ['2026-03-26 09:19:10', '2026-03-25 06:44:44'],
            'Ip_Lat': [-11.07, -12.97], 'Ip_Lon': [-37.33, -38.49],
            'Ip_Cidade': ['A', 'B'], 'Ip_Pais': ['Brazil'] * 2,
            'Ip_Dono': ['X', 'Y'],
            'Ip_Proxy': [False] * 2, 'Ip_Hospedagem': [False] * 2,
            'Ip_Movel': [False, True],
        })
        html = generate_html_report(df, 'alvo')
        texto = html.decode('utf-8') if isinstance(html, bytes) else html
        self.assertNotIn('cartocdn', texto)
        self.assertNotIn('cartodb', texto.lower())
        # e o fundo tem de continuar existindo
        self.assertIn('tile.openstreetmap.org', texto)

    def test_todos_os_estilos_produzem_fundo(self):
        import folium
        from helpers.geo import TILE_SOURCES, add_base_layer

        for nome in TILE_SOURCES:
            with self.subTest(estilo=nome):
                fmap = folium.Map(location=[-11, -37], tiles=None)
                self.assertEqual(add_base_layer(fmap, nome), nome)
                render = fmap.get_root().render()
                self.assertNotIn('cartocdn', render)
                self.assertIn(TILE_SOURCES[nome]['url'].split('{')[0], render)

    def test_estilo_desconhecido_cai_no_padrao(self):
        """Um estilo gravado na sessão que não exista mais não pode deixar o
        mapa sem fundo nenhum."""
        import folium
        from helpers.geo import DEFAULT_TILE, add_base_layer

        fmap = folium.Map(location=[-11, -37], tiles=None)
        self.assertEqual(add_base_layer(fmap, 'Voyager'), DEFAULT_TILE)
        self.assertIn('arcgisonline', fmap.get_root().render())

    def test_gabarito_do_arcgis_usa_z_y_x(self):
        """O ArcGIS serve linha antes de coluna. Trocar a ordem devolve tiles
        de outro lugar do mundo, com HTTP 200 e sem erro nenhum."""
        from helpers.geo import TILE_SOURCES

        for nome, fonte in TILE_SOURCES.items():
            if 'arcgisonline' not in fonte['url']:
                continue
            with self.subTest(estilo=nome):
                self.assertTrue(fonte['url'].endswith('/{z}/{y}/{x}'), fonte['url'])


class TestExportLimits(unittest.TestCase):
    def test_xlsx_divide_em_abas_acima_do_limite_da_planilha(self):
        """O limite de 1.048.576 linhas é do formato XLSX, não da ferramenta.
        Acima dele o resultado continua em abas — nada é omitido."""
        import io as _io
        from openpyxl import load_workbook
        from file_handler import export_xlsx_colored
        df = pd.DataFrame({'a': range(25), 'b': ['x'] * 25})
        buf = _io.BytesIO()
        # sheet_row_limit reduzido: exercita a divisão sem gerar um milhão de linhas
        resumo = export_xlsx_colored(df, buf, sheet_row_limit=10)
        self.assertEqual(resumo, {'linhas': 25, 'abas': 3})

        buf.seek(0)
        wb = load_workbook(buf)
        self.assertEqual(wb.sheetnames,
                         ['Resultado', 'Resultado (2)', 'Resultado (3)'])
        linhas = 0
        for nome in wb.sheetnames:
            ws = wb[nome]
            # cabeçalho repetido em cada aba, senão a continuação chega sem colunas
            self.assertEqual([c.value for c in ws[1]], ['a', 'b'])
            linhas += ws.max_row - 1
        self.assertEqual(linhas, len(df))

    def test_xlsx_rejects_oversized_frame(self):
        """Acima do teto de abas a mensagem precisa ser acionável e apontar o
        CSV, não um MemoryError no meio de uma execução de horas."""
        import io as _io
        from file_handler import export_xlsx_colored, XLSX_MAX_SHEETS
        df = pd.DataFrame({'a': range(XLSX_MAX_SHEETS + 1)})
        with self.assertRaises(ValueError) as ctx:
            export_xlsx_colored(df, _io.BytesIO(), sheet_row_limit=1)
        self.assertIn('linhas', str(ctx.exception))
        self.assertIn('CSV', str(ctx.exception))

    def test_xlsx_em_disco_e_atomico(self):
        """Truncar-e-escrever já destruiu artefato em falha no meio da gravação.
        Numa falha, nenhum `.xlsx` com cara de export completo pode sobrar."""
        import os
        import file_handler
        from file_handler import export_xlsx_to_disk

        chamadas = {}

        def _falha_no_meio(df, output, **kwargs):
            chamadas['parcial'] = output
            open(output, 'wb').write(b'lixo parcial')
            raise MemoryError('falha simulada no meio da gravação')

        original = file_handler.export_xlsx_colored
        file_handler.export_xlsx_colored = _falha_no_meio
        try:
            with self.assertRaises(MemoryError):
                export_xlsx_to_disk(pd.DataFrame({'a': [1]}), 'atomico_teste')
        finally:
            file_handler.export_xlsx_colored = original

        parcial = chamadas['parcial']
        self.assertTrue(parcial.endswith('.part'))
        self.assertFalse(os.path.exists(parcial), 'parcial não foi removido')
        self.assertFalse(os.path.exists(parcial[:-len('.part')]),
                         'destino final não pode existir após falha')

    def test_xlsx_em_disco_grava_e_declara_o_resumo(self):
        import os
        from openpyxl import load_workbook
        from file_handler import export_xlsx_to_disk

        df = pd.DataFrame({'Ip': ['1.1.1.1'] * 5,
                           'Reputação': ['🟢 Residencial'] * 5})
        caminho, resumo = export_xlsx_to_disk(df, 'disco_teste')
        try:
            self.assertTrue(os.path.exists(caminho))
            self.assertEqual(resumo, {'linhas': 5, 'abas': 1})
            ws = load_workbook(caminho).active
            self.assertEqual(ws.max_row, 6)
            self.assertIn('DCFCE7', ws.cell(2, 1).fill.start_color.rgb)
        finally:
            os.remove(caminho)

    def test_xlsx_neutraliza_formula_injetada(self):
        """O caminho da interface exportava sem sanitizar — só o do pipeline
        passava por `sanitize_dataframe_for_csv`. Uma célula iniciada por `=`,
        `+`, `-` ou `@` vira fórmula ao abrir a planilha, então a sanitização
        mora dentro do export e não a cargo de quem chama."""
        import io as _io
        from openpyxl import load_workbook
        from file_handler import export_xlsx_colored

        payloads = ["=cmd|' /C calc'!A0", '+1+1', '@SUM(1+1)', '-2+3']
        df = pd.DataFrame({'Ip_Dono': payloads,
                           'Reputação': ['🟢 Residencial'] * len(payloads)})
        original = df['Ip_Dono'].tolist()

        buf = _io.BytesIO()
        export_xlsx_colored(df, buf)
        buf.seek(0)
        ws = load_workbook(buf).active
        for linha, payload in enumerate(payloads, 2):
            with self.subTest(payload=payload):
                valor = ws.cell(linha, 1).value
                self.assertTrue(valor.startswith("'"), valor)
                self.assertEqual(valor[1:], payload)

        # O frame do chamador não pode ser alterado pelo export.
        self.assertEqual(df['Ip_Dono'].tolist(), original)
        # E a cor por reputação continua saindo.
        self.assertIn('DCFCE7', ws.cell(2, 1).fill.start_color.rgb)

    def test_xlsx_preserves_reputation_colors(self):
        import io as _io
        from openpyxl import load_workbook
        from file_handler import export_xlsx_colored
        df = pd.DataFrame({
            'Ip': ['1.1.1.1', '2.2.2.2'],
            'Reputação': ['🔴 Proxy / VPN / Tor', '🟢 Residencial'],
        })
        buf = _io.BytesIO()
        export_xlsx_colored(df, buf)
        buf.seek(0)
        ws = load_workbook(buf).active
        self.assertEqual(ws.cell(1, 1).value, 'Ip')
        self.assertIn('FECACA', ws.cell(2, 1).fill.start_color.rgb)
        self.assertIn('DCFCE7', ws.cell(3, 1).fill.start_color.rgb)


class TestIntegrityHash(unittest.TestCase):
    def test_chunked_hash_matches_single_pass(self):
        """O hash entra na cadeia de custódia de relatórios já entregues:
        a versão em blocos precisa ser byte-a-byte idêntica à anterior."""
        import hashlib
        from html_report_generator import _hash_dataframe_csv
        for n, chunk in ((0, 50), (1, 50), (999, 100), (5000, 512)):
            df = pd.DataFrame({'Ip': ['1.1.1.%d' % (i % 256) for i in range(n)],
                               'X': range(n)})
            esperado = hashlib.sha256(df.to_csv(index=False).encode('utf-8')).hexdigest()
            self.assertEqual(_hash_dataframe_csv(df, chunk_rows=chunk), esperado,
                             'n=%d chunk=%d' % (n, chunk))


class TestSalvarExportacao(unittest.TestCase):
    """Tail de exportação compartilhado pelos dois pipelines."""

    def setUp(self):
        import tempfile
        self.dir = tempfile.mkdtemp()

    def _df(self, n=2):
        return pd.DataFrame({
            'Ip': ['1.1.1.1'] * n,
            'Data': ['2025-01-01 10:00:00'] * n,
            'Reputação': ['🟢 Residencial'] * n,
            'Ignorada': ['x'] * n,
        })

    def test_gera_csv_e_xlsx(self):
        import os
        from file_handler import salvar_exportacao
        out = os.path.join(self.dir, 'r.csv')
        df_export, xlsx = salvar_exportacao(
            self._df(), ['Ip', 'Data', 'Reputação'], out)
        self.assertTrue(os.path.exists(out))
        self.assertTrue(xlsx and os.path.exists(xlsx))
        # colunas fora da lista de exportação não vazam para o artefato
        self.assertNotIn('Ignorada', df_export.columns)

    def test_falha_no_xlsx_nao_aborta_o_processamento(self):
        """Regressão: uma falha no Excel não pode derrubar uma execução inteira.
        O CSV é o artefato primário e precisa ser gravado mesmo assim."""
        import os
        import file_handler
        from file_handler import salvar_exportacao

        def _falha(*args, **kwargs):
            raise ValueError('teto simulado')

        out = os.path.join(self.dir, 'grande.csv')
        original = file_handler.export_xlsx_colored
        file_handler.export_xlsx_colored = _falha
        try:
            df_export, xlsx = salvar_exportacao(self._df(3), ['Ip', 'Data'], out)
        finally:
            file_handler.export_xlsx_colored = original

        self.assertTrue(os.path.exists(out))
        self.assertEqual(len(df_export), 3)
        self.assertIsNone(xlsx)

    def test_colunas_ausentes_sao_ignoradas(self):
        import os
        from file_handler import salvar_exportacao
        out = os.path.join(self.dir, 'p.csv')
        df_export, _ = salvar_exportacao(
            self._df(), ['Ip', 'NaoExiste', 'Data'], out)
        self.assertEqual(list(df_export.columns), ['Ip', 'Data'])
