import numpy as np
import pandas as pd
import logging
import os
import re
import json
from pathlib import Path
from data_processor import (COLUNAS_MODELO, FAMILIAS_GOOGLE_AGREGADO, cabecalho_csv,
                          colunas_export_para, familia_formato, formato_google_agregado,
                          detectar_separador_csv, extrair_ips_de_texto, processar_resultados,
                          extrair_ips_do_formato_preservation_google,
                          extrair_ips_do_formato_google_atividades,
                          extrair_ips_do_formato_google_conexoes,
                          TZ_LABEL, get_periodo, format_iso_date)
from api_client import is_valid_ip


class FormatoSaidaIncompativel(ValueError):
    """O arquivo de saída incremental guarda o resultado de outro formato.

    Os dois CSVs "Aggregated Activities" do Google são respostas distintas:
    misturá-los (ou misturar um deles com acessos de outra plataforma) no
    mesmo arquivo produziria uma tabela em que metade das linhas é intervalo
    e metade é par de identificadores. A recusa acontece antes de consultar a
    API e sem tocar no arquivo existente.
    """

    def __init__(self, arquivo, existente, novo):
        self.arquivo = arquivo
        self.existente = existente
        self.novo = novo
        super().__init__(
            f"O arquivo de saída {arquivo} já contém resultado do formato "
            f"'{existente}', e o novo é '{novo}'. Nada foi misturado nem "
            "sobrescrito: escolha outro arquivo de saída.")


def _recusar_mistura(df_novo, df_existente, output_file):
    """Levanta `FormatoSaidaIncompativel` se um dos lados é Google agregado
    e os formatos diferem."""
    novo, existente = familia_formato(df_novo), familia_formato(df_existente)
    if novo != existente and (novo in FAMILIAS_GOOGLE_AGREGADO
                              or existente in FAMILIAS_GOOGLE_AGREGADO):
        raise FormatoSaidaIncompativel(output_file, existente, novo)


_PARSER_GOOGLE_AGREGADO = {
    'google_atividades': extrair_ips_do_formato_google_atividades,
    'google_conexoes': extrair_ips_do_formato_google_conexoes,
}


def _ler_texto(caminho):
    """Texto do arquivo: UTF-8 (com ou sem BOM), senão cp1252, senão latin-1."""
    with open(caminho, 'rb') as f:
        dados = f.read()
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            return dados.decode(encoding)
        except UnicodeDecodeError:
            continue
    return dados.decode('latin-1')


def _primeira_linha_arquivo(caminho):
    """Primeira linha não vazia, sem ler o arquivo inteiro."""
    with open(caminho, 'r', encoding='utf-8-sig', errors='replace') as f:
        for linha in f:
            if linha.strip():
                return linha
    return ''

# Limite do próprio formato XLSX: 1.048.576 linhas por planilha, menos o
# cabeçalho. Acima disso o resultado continua em abas adicionais.
XLSX_SHEET_ROWS = 1_048_575
# Teto de abas de continuação (~10,4 milhões de linhas). Acima disso o tempo de
# geração — cerca de um minuto por 100 mil linhas — deixa de fazer sentido e o
# CSV é o artefato apropriado.
XLSX_MAX_SHEETS = 10

# Configurar logger para este módulo
logger = logging.getLogger(__name__)


def _fmt_int(n):
    """Formata inteiro no padrão pt-BR: 202128 -> '202.128'."""
    try:
        return '{:,}'.format(int(n)).replace(',', '.')
    except (TypeError, ValueError):
        return str(n)


def _parse_datas(series):
    """Converte a coluna `Data` bruta para datetime ingênuo (hora local do log).

    Substitui dois atalhos que liam datas brasileiras como mês/dia: escolher o
    formato olhando só a primeira linha, e `pd.to_datetime(valor)` linha a
    linha (`05/03/2024` virava 3 de maio). `parse_data` infere a ordem pela
    coluna inteira. Fuso explícito é descartado mantendo a hora de parede,
    como o `strftime` anterior fazia.
    """
    from validators import parse_data
    dt = parse_data(series)
    if not pd.api.types.is_datetime64_any_dtype(dt):
        # Fusos mistos: pandas devolve objetos; normaliza um a um.
        dt = pd.to_datetime(
            dt.map(lambda x: x.replace(tzinfo=None)
                   if pd.notna(x) and getattr(x, 'tzinfo', None) else x),
            errors='coerce')
    elif getattr(dt.dt, 'tz', None) is not None:
        dt = dt.dt.tz_localize(None)
    return dt


def _preencher_periodo(df, col='Periodo'):
    """Deriva `Periodo` de `Data`; deixa vazio quando não dá para saber.

    Gravar '☀️ Diurno' em registros sem data parseável fabrica um dado que
    o documento não afirma — num laudo, ausência declarada é melhor que
    valor inventado.
    """
    if 'Data' not in df.columns:
        df[col] = None
        return
    try:
        data_dt = _parse_datas(df['Data'])
        df[col] = data_dt.dt.hour.map(lambda h: get_periodo(int(h)) if pd.notna(h) else None)
    except (AttributeError, TypeError, ValueError) as e:
        logger.warning(f"Erro ao calcular Periodo: {e}")
        df[col] = None


def carregar_arquivo_log(input_file, update_callback=None, alvo='desconhecido'):
    """Carrega um arquivo de log e extrai os IPs com seus dados temporais"""
    ext = Path(input_file).suffix.lower()

    try:
        if update_callback:
            update_callback(f"Carregando arquivo {input_file}")

        # Extrair data do nome do arquivo se possível
        filename = os.path.basename(input_file)
        data_arquivo = None

        # Verificar formato específico 20250310123909
        timestamp_match = re.search(r'(\d{14})', filename)
        if timestamp_match:
            timestamp_str = timestamp_match.group(1)
            try:
                data_arquivo = f"{timestamp_str[0:4]}-{timestamp_str[4:6]}-{timestamp_str[6:8]} {timestamp_str[8:10]}:{timestamp_str[10:12]}:{timestamp_str[12:14]}"
                logger.info(f"Data extraída do nome do arquivo: {data_arquivo}")
            except (TypeError, ValueError):
                pass

        if ext == '.csv':
            # Google Aggregated Activities: reconhecido pelo cabeçalho, antes da
            # leitura com inferência de tipos — o GAIA ID (até 21 dígitos)
            # estouraria o int64, e o laço genérico abaixo pegaria só o último
            # uso, em UTC rotulado como fuso local.
            agregado = formato_google_agregado(cabecalho_csv(_primeira_linha_arquivo(input_file)))
            if agregado:
                if update_callback:
                    update_callback(f"Formato {agregado} detectado")
                return _PARSER_GOOGLE_AGREGADO[agregado](
                    _ler_texto(input_file), update_callback, alvo=alvo)

            # Detectar separador
            separador = detectar_separador_csv(input_file)
            logger.info(f"Separador detectado: '{separador}'")

            # Tentar diferentes encodings
            for encoding in ['utf-8', 'latin1', 'cp1252']:
                try:
                    df = pd.read_csv(input_file, sep=separador, encoding=encoding)
                    break
                except UnicodeDecodeError:
                    continue
                except Exception as e:
                    raise e
            else:
                df = pd.read_csv(input_file, sep=separador, encoding='utf-8', encoding_errors='replace')
        elif ext in ('.xls', '.xlsx'):
            df = pd.read_excel(input_file)
            agregado = formato_google_agregado(df.columns)
            if agregado:
                # CSV do Google regravado no Excel: relido como texto.
                return _PARSER_GOOGLE_AGREGADO[agregado](
                    pd.read_excel(input_file, dtype=str), update_callback, alvo=alvo)
        elif ext == '.txt':
            # Usar função específica para extrair IPs de texto
            return extrair_ips_de_texto(input_file, True, update_callback, alvo=alvo)
        elif ext in ('.html', '.htm'):
            # HTML de plataformas (WhatsApp, Meta, Google)
            from html_parser import parse_html_file
            with open(input_file, 'r', encoding='utf-8', errors='replace') as f:
                html_content = f.read()
            df_html, platform = parse_html_file(html_content, alvo=alvo, update_callback=update_callback)
            if update_callback:
                update_callback(f"HTML ({platform}): {len(df_html)} registros extraídos")
            return df_html
        elif ext == '.pdf':
            # Extrair texto do PDF e processar como texto
            try:
                import pdfplumber
                text_parts = []
                with pdfplumber.open(input_file) as pdf:
                    for page_num, page in enumerate(pdf.pages, 1):
                        try:
                            page_text = page.extract_text()
                            if page_text:
                                text_parts.append(page_text)
                        except Exception as e:
                            logger.warning("Página %d do PDF não extraída: %s", page_num, e)
                content = '\n'.join(text_parts)
                if not content.strip():
                    raise ValueError("PDF sem conteúdo de texto extraível")
                if update_callback:
                    update_callback(f"Texto extraído do PDF ({len(text_parts)} páginas)")
                return extrair_ips_de_texto(content, False, update_callback, alvo=alvo)
            except ImportError:
                raise ValueError("Biblioteca 'pdfplumber' necessária para PDFs. Instale com: pip install pdfplumber")
        else:
            raise ValueError(f"Formato de arquivo não suportado: {ext}")

        # Detectar formato Preservation Google (CSV do Google Takeout)
        preservation_google_cols = ['Activity Timestamp', 'IP Address', 'User Agent String']
        if all(col in df.columns for col in preservation_google_cols):
            if update_callback:
                update_callback("Formato Preservation Google detectado")
            return extrair_ips_do_formato_preservation_google(df, update_callback, alvo=alvo)

        # Se já temos um arquivo no formato do modelo
        modelo_columns_lower = [col.lower() for col in COLUNAS_MODELO]
        df_columns_lower = [col.lower() for col in df.columns]

        if all(col in df_columns_lower for col in ['ip', 'data']):
            # Padronizar nomes das colunas para o modelo
            rename_map = {}
            for i, col in enumerate(df.columns):
                col_lower = col.lower()
                if col_lower in modelo_columns_lower:
                    index = modelo_columns_lower.index(col_lower)
                    if col != COLUNAS_MODELO[index]:
                        rename_map[col] = COLUNAS_MODELO[index]

            if rename_map:
                df = df.rename(columns=rename_map)

            # Garantir que todas as colunas do modelo existam
            for col in COLUNAS_MODELO:
                if col not in df.columns:
                    if col == 'Alvo':
                        df[col] = 'desconhecido'
                    elif col == 'Data_Fuso':
                        df[col] = TZ_LABEL
                    elif col in ['Ip_Movel', 'Ip_Proxy', 'Ip_Hospedagem', 'Ip_Tor']:
                        df[col] = False
                    elif col == 'Periodo':
                        _preencher_periodo(df)
                    elif col == 'ISO_Date':
                        # Gerar ISO_Date a partir da Data
                        if 'Data' in df.columns:
                            try:
                                data_dt = _parse_datas(df['Data'])
                                df['ISO_Date'] = data_dt.apply(
                                    lambda dt: format_iso_date(dt.to_pydatetime()) if pd.notna(dt) else None
                                )
                            except (AttributeError, TypeError, ValueError) as e:
                                logger.warning(f"Erro ao calcular ISO_Date: {e}")
                                df['ISO_Date'] = None
                        else:
                            df['ISO_Date'] = None
                    else:
                        df[col] = None

            # Garantir a ordem exata das colunas do modelo (preservar colunas extras por formato)
            if 'User_ID' in df.columns:
                colunas = [col for col in COLUNAS_MODELO]
                idx_ip = colunas.index('Ip')
                for i, extra_col in enumerate(['User_ID', 'Username', 'Email']):
                    colunas.insert(idx_ip + 1 + i, extra_col)
                df = df[[c for c in colunas if c in df.columns]]
            elif 'Porta' in df.columns:
                colunas = [col for col in COLUNAS_MODELO]
                idx_ip = colunas.index('Ip')
                colunas.insert(idx_ip + 1, 'Porta')
                df = df[[c for c in colunas if c in df.columns]]
            else:
                df = df[COLUNAS_MODELO]

            # Filtrar apenas linhas com IPs válidos
            df['Ip_Valido'] = df['Ip'].apply(lambda x: is_valid_ip(str(x)))
            df = df[df['Ip_Valido']].copy()
            df = df.drop('Ip_Valido', axis=1)

            # Definir Alvo
            df['Alvo'] = alvo

            return df
        else:
            # Criar DataFrame com estrutura do modelo a partir dos dados
            registros = []

            for _, row in df.iterrows():
                # Extrair IP e data
                ip = None
                data = None

                # Tentar encontrar o IP e a data nas colunas
                for col in df.columns:
                    col_str = str(col).lower()
                    valor = row[col]

                    # Procurar IP
                    if not ip and ('ip' in col_str or 'address' in col_str or 'endereco' in col_str):
                        if is_valid_ip(str(valor)):
                            ip = str(valor)

                    # Procurar data
                    if not data and ('data' in col_str or 'date' in col_str or 'time' in col_str or 'hora' in col_str):
                        try:
                            data = str(valor)
                        except (TypeError, ValueError):
                            pass

                # Se não achou IP, tentar qualquer coluna
                if not ip:
                    for col in df.columns:
                        if is_valid_ip(str(row[col])):
                            ip = str(row[col])
                            break

                # Se tem IP válido, adicionar ao resultado
                if ip and is_valid_ip(ip):
                    # Se não tem data, usar a extraída do nome do arquivo (ou deixar vazio)
                    if not data:
                        data = data_arquivo

                    # Data, Periodo e ISO_Date são derivados depois do laço,
                    # numa passada vetorizada sobre a coluna inteira: a ordem
                    # dia/mês só se decide olhando todas as datas.
                    periodo = None
                    iso_date = None

                    # Criar registro no formato do modelo
                    registro = {
                        'Alvo': alvo,
                        'Ip': ip,
                        'Data': data,
                        'Data_Fuso': TZ_LABEL if data else None,
                        'Ip_Dono': None,
                        'Ip_AS': None,
                        'Ip_Regiao': None,
                        'Ip_Cidade': None,
                        'Ip_Pais': None,
                        'Ip_Pais_Codigo': None,
                        'Ip_Movel': False,
                        'Ip_Proxy': False,
                        'Ip_Hospedagem': False,
                        'Ip_Tor': False,
                        'Ip_Lat': None,
                        'Ip_Lon': None,
                        'Periodo': periodo,
                        'ISO_Date': iso_date
                    }

                    registros.append(registro)

            # Criar DataFrame com registros extraídos
            df_resultado = pd.DataFrame(registros)

            # Se vazio, criar DataFrame vazio com colunas do modelo
            if df_resultado.empty:
                df_resultado = pd.DataFrame(columns=COLUNAS_MODELO)
            else:
                data_dt = _parse_datas(df_resultado['Data'])
                ok = data_dt.notna()
                falhas = int((~ok & df_resultado['Data'].notna()).sum())
                if falhas:
                    # Valor não interpretável fica como veio do documento, sem
                    # período nem ISO_Date — nada é inventado.
                    logger.warning("%d datas não interpretáveis mantidas como no original", falhas)
                df_resultado.loc[ok, 'Data'] = data_dt[ok].dt.strftime('%Y-%m-%d %H:%M:%S')
                df_resultado.loc[ok, 'Periodo'] = data_dt[ok].dt.hour.map(get_periodo)
                df_resultado.loc[ok, 'ISO_Date'] = data_dt[ok].map(
                    lambda dt: format_iso_date(dt.to_pydatetime()))
                # Garantir a ordem das colunas
                for col in COLUNAS_MODELO:
                    if col not in df_resultado.columns:
                        if col in ['Ip_Movel', 'Ip_Proxy', 'Ip_Hospedagem', 'Ip_Tor']:
                            df_resultado[col] = False
                        elif col == 'Alvo':
                            df_resultado[col] = 'desconhecido'
                        elif col == 'Data_Fuso':
                            df_resultado[col] = TZ_LABEL
                        elif col == 'Periodo':
                            _preencher_periodo(df_resultado)
                        else:
                            df_resultado[col] = None

                df_resultado = df_resultado[COLUNAS_MODELO]

            return df_resultado

    except Exception as e:
        raise e


async def processar_log_acesso_async(input_file_or_content, output_file, is_file=True, batch_size=500, period=0,
                               cache_file=None, incremental=True,
                               update_callback=None, progress_callback=None, alvo='desconhecido', api_key=None):
    """Processa arquivo de log ou texto, consultando IPs na API"""
    if cache_file is None:
        from api_client import IP_CACHE_PATH
        cache_file = IP_CACHE_PATH
    try:
        # Atualizar status
        if update_callback:
            if is_file:
                update_callback(f"Carregando arquivo: {input_file_or_content}")
            else:
                update_callback("Processando texto colado")

        # Carregar os dados
        if is_file:
            df = carregar_arquivo_log(input_file_or_content, update_callback, alvo=alvo)
        else:
            df = extrair_ips_de_texto(input_file_or_content, False, update_callback, alvo=alvo)

        # O formato (e com ele as colunas de export) sai das colunas extras
        # que o parser deixou: Porta, Evento, User_Agent, User_ID, Vinculado_*...
        if df.empty:
            if update_callback:
                update_callback("Nenhum IP válido encontrado nos dados")
            return pd.DataFrame(columns=colunas_export_para(df))

        total_records = len(df)
        total_unique_all = df['Ip'].nunique()
        if update_callback:
            update_callback(f"Extraídos {total_records} registros com {total_unique_all} IPs únicos")

        # Verificar se o arquivo já existe e carregar para modo incremental
        df_resultado = None
        if incremental and os.path.exists(output_file):
            try:
                ext_saida = Path(output_file).suffix.lower()
                if ext_saida in ('.xlsx', '.xls'):
                    df_resultado = pd.read_excel(output_file)
                else:
                    sep_existente = detectar_separador_csv(output_file)
                    for encoding in ['utf-8', 'latin1', 'cp1252']:
                        try:
                            df_resultado = pd.read_csv(output_file, sep=sep_existente, encoding=encoding)
                            break
                        except UnicodeDecodeError:
                            continue
                        except (OSError, ValueError, pd.errors.ParserError) as e:
                            raise e
                    else:
                        df_resultado = pd.read_csv(output_file, sep=sep_existente, encoding='utf-8', encoding_errors='replace')

                # Garantir que temos as colunas do modelo
                for col in COLUNAS_MODELO:
                    if col not in df_resultado.columns:
                        if col in ['Ip_Movel', 'Ip_Proxy', 'Ip_Hospedagem', 'Ip_Tor']:
                            df_resultado[col] = False
                        elif col == 'Alvo':
                            df_resultado[col] = alvo
                        elif col == 'Data_Fuso':
                            df_resultado[col] = TZ_LABEL
                        elif col == 'Periodo':
                            _preencher_periodo(df_resultado)
                        else:
                            df_resultado[col] = None

                # Recuperar Lat/Lon do cache se ausentes no CSV
                lat_missing = df_resultado['Ip_Lat'].isna().all() if 'Ip_Lat' in df_resultado.columns else True
                if lat_missing and cache_file and os.path.exists(cache_file):
                    try:
                        # Leitor único: desembrulha o envelope assinado e só
                        # devolve coordenadas dentro do TTL — geolocalização
                        # expirada não é servida como atual.
                        from api_client import filter_fresh_cache, read_ip_cache
                        ip_cache_data, expiradas = filter_fresh_cache(read_ip_cache(cache_file))
                        df_resultado['Ip_Lat'] = df_resultado['Ip'].map(
                            lambda ip: ip_cache_data.get(ip, {}).get('Ip_Lat'))
                        df_resultado['Ip_Lon'] = df_resultado['Ip'].map(
                            lambda ip: ip_cache_data.get(ip, {}).get('Ip_Lon'))
                        recovered = df_resultado['Ip_Lat'].notna().sum()
                        if update_callback and recovered > 0:
                            update_callback(f"Coordenadas recuperadas do cache para {recovered} registros")
                        if update_callback and expiradas:
                            update_callback(f"{expiradas} entradas do cache ignoradas por estarem "
                                            "expiradas ou sem data")
                    except (OSError, UnicodeError, json.JSONDecodeError, TypeError, ValueError) as e:
                        logger.exception("Erro ao recuperar coordenadas do cache")
                        if update_callback:
                            update_callback(f"⚠️ Coordenadas não recuperadas do cache: {e}")

                if update_callback:
                    update_callback(f"Arquivo existente carregado: {len(df_resultado)} registros")
            except Exception as e:
                if update_callback:
                    update_callback(f"Erro ao carregar arquivo existente: {e}")
                logger.warning(f"Erro ao carregar arquivo existente: {e}")

        # Fora do try acima: a recusa tem de chegar à página, não virar aviso.
        if df_resultado is not None:
            _recusar_mistura(df, df_resultado, output_file)

        # Obter lista de IPs únicos para consulta
        ips_unicos = df['Ip'].unique()

        # Verificar se já temos informações sobre algum IP
        if df_resultado is not None:
            # Só considerar IP como "processado" se tiver dados de enriquecimento válidos
            df_enriquecidos = df_resultado[_tem_enriquecimento(df_resultado)]
            ips_processados = set(df_enriquecidos['Ip'].astype(str))
            ips_novos = [ip for ip in ips_unicos if ip not in ips_processados]
            enriquecimento_existente = _enriquecimento_por_ip(df_enriquecidos, update_callback)

            if update_callback:
                update_callback(f"{len(ips_unicos) - len(ips_novos)} IPs já enriquecidos, {len(ips_novos)} IPs para consultar")

            ips_unicos = ips_novos

        # Atualizar progresso inicial
        total_ips = len(ips_unicos)
        if progress_callback:
            progress_callback(0, total_ips)

        resultados_ips = {}
        if total_ips == 0:
            # Sem IP novo não há consulta, mas pode haver registro novo de IP
            # já enriquecido (outra data, outra sessão), e ele segue para a
            # mescla. Devolver o arquivo existente aqui descartava esses
            # registros em silêncio.
            if update_callback:
                update_callback("Nenhum IP novo para consultar")
        else:
            # Enriquecimento via serviço compartilhado (mesmo caminho de interception_parser)
            from enrich_service import enrich_ips

            def _on_batch_progress(processed, total):
                if progress_callback:
                    progress_callback(min(processed, total_ips), total_ips, total_records)

            if update_callback:
                update_callback(f"Enviando {total_ips} IPs para o endpoint batch...")

            resultados_ips, _client = await enrich_ips(
                list(ips_unicos),
                cache_file=cache_file,
                api_key=api_key,
                batch_size=batch_size,
                period=period,
                update_callback=update_callback,
                progress_callback=_on_batch_progress,
                save_cache=True,
            )

        # Processar resultados finais
        df_processado = processar_resultados(df, resultados_ips)

        # Combinar com dados existentes
        if df_resultado is not None:
            # A Reputação relida do CSV só cobre as linhas antigas, e
            # add_reputacao_column não recalcula coluna que já existe: as
            # novas saíam sem rótulo. Ela deriva das colunas do IP e é
            # recalculada para todas.
            df_existente = df_resultado.drop(columns='Reputação', errors='ignore')
            # IP já enriquecido não volta à API: o enriquecimento das linhas
            # dele vem do próprio arquivo que está sendo estendido. Antes do
            # dedup, para que a linha nova preenchida vença a antiga.
            preenchidas = np.concatenate([
                _preencher_enriquecimento(df_existente, enriquecimento_existente),
                _preencher_enriquecimento(df_processado, enriquecimento_existente)])
            df_final = pd.concat([df_existente, df_processado], ignore_index=True)
            dedup_cols = _incremental_dedup_columns(df_final)
            # Chave comparada como texto: o CSV relido infere `349` como
            # número e o registro novo traz `'349'` — sem isso o mesmo registro
            # entraria duas vezes. Vazio (None/NaN) conta como igual a vazio.
            chave = df_final[dedup_cols].fillna('').astype(str)
            # Na colisão vence a linha mais recente, mas uma linha enriquecida
            # nunca é trocada por outra sem enriquecimento: as enriquecidas
            # vão para o fim da ordem (estável) em que o keep='last' decide.
            ordem = np.argsort(_tem_enriquecimento(df_final).to_numpy(), kind='stable')
            repetidas = np.empty(len(df_final), dtype=bool)
            repetidas[ordem] = chave.iloc[ordem].duplicated(keep='last').to_numpy()
            df_final = df_final[~repetidas]
            preenchidas = preenchidas[~repetidas]

            existentes = df_final.index < len(df_existente)
            n_do_log = int((preenchidas & ~existentes).sum())
            n_reparadas = int((preenchidas & existentes).sum())
            df_final.index = pd.RangeIndex(len(df_final))
            if update_callback and n_do_log:
                update_callback(f"{_fmt_int(n_do_log)} registros do log processado, de IPs já "
                                "enriquecidos, receberam o enriquecimento do arquivo existente")
            if n_reparadas:
                # O arquivo entregue muda de conteúdo: isso se declara.
                logger.warning("%d registros do arquivo existente estavam sem enriquecimento "
                               "de IP enriquecido em outras linhas; preenchidos", n_reparadas)
                if update_callback:
                    update_callback(f"⚠️ {_fmt_int(n_reparadas)} registros do arquivo existente "
                                    "estavam sem enriquecimento, embora o mesmo IP esteja "
                                    "enriquecido em outras linhas dele, e foram preenchidos")
        else:
            df_final = df_processado

        # Normalizar nomes de provedores (Claro S.A/NXT/Net → Claro, etc.)
        from data_processor import normalizar_provedor_df
        df_final = normalizar_provedor_df(df_final)

        # Adicionar coluna Reputação
        df_final = _add_reputacao_column(df_final)

        # Colunas de exportação do formato (extras de cada provedor logo após Ip)
        export_cols = [c for c in colunas_export_para(df_final) if c in df_final.columns]
        salvar_exportacao(df_final, export_cols, output_file,
                          update_callback, gerar_xlsx=False)

        return df_final

    except Exception as e:
        logger.error(f"Erro ao processar log de acesso: {str(e)}", exc_info=True)
        if update_callback:
            update_callback(f"Erro: {str(e)}")
        raise


# Campos que vêm da consulta do IP, não do registro do log.
_CAMPOS_DO_IP = [c for c in COLUNAS_MODELO if c.startswith('Ip_')]


def _tem_enriquecimento(df):
    """Linhas com enriquecimento: o mesmo critério que dispensa a consulta."""
    return df['Ip_Dono'].notna() & (df['Ip_Dono'] != '')


def _enriquecimento_por_ip(df_enriquecidos, update_callback=None):
    """Enriquecimento de cada IP já consultado, uma linha por IP (índice `Ip`).

    IP com mais de uma versão de enriquecimento no arquivo (editado à mão,
    juntado de outra saída) fica de fora, e isso é declarado: escolher uma
    delas afirmaria o que o arquivo não afirma.
    """
    campos = [c for c in _CAMPOS_DO_IP if c in df_enriquecidos.columns]
    tabela = (df_enriquecidos[['Ip', *campos]]
              .assign(Ip=lambda d: d['Ip'].astype(str))
              .drop_duplicates())
    divergentes = tabela['Ip'].duplicated(keep=False)
    if divergentes.any():
        n = tabela.loc[divergentes, 'Ip'].nunique()
        logger.warning("%d IPs com mais de um enriquecimento no arquivo existente", n)
        if update_callback:
            update_callback(f"⚠️ {_fmt_int(n)} IPs têm mais de um enriquecimento no arquivo "
                            "existente: registros novos deles ficam sem enriquecimento")
    return tabela[~divergentes].set_index('Ip')


def _preencher_enriquecimento(df, enriquecimento):
    """Preenche, no lugar, as linhas sem enriquecimento de IPs de `enriquecimento`.

    Sobrescreve todos os campos do IP, não só os vazios: a linha sem
    enriquecimento traz `False` nos booleanos (default do parser), e um
    `fillna` deixaria um proxy marcado como não proxy.

    Returns:
        máscara (ndarray) das linhas preenchidas.
    """
    ip = df['Ip'].astype(str)
    alvo = (~_tem_enriquecimento(df) & ip.isin(enriquecimento.index)).to_numpy()
    if alvo.any():
        for campo in enriquecimento.columns:
            df[campo] = ip.map(enriquecimento[campo]).where(alvo, df[campo])
    return alvo


def _incremental_dedup_columns(df):
    """Colunas de dedup incremental: Ip+Data e chaves de formato.

    `Vinculado_*` é chave nas conexões do Google: o mesmo IP com o mesmo
    último uso aparece uma vez por identificador vinculado, e sem elas um
    reprocessamento colapsaria esses pares num só.
    """
    cols = []
    for c in ('Ip', 'Data', 'Porta', 'Evento', 'User_ID', 'User_Agent',
              'Vinculado_Tipo', 'Vinculado_Valor'):
        if c in df.columns:
            cols.append(c)
    return cols or list(df.columns[:1])


def _add_reputacao_column(df):
    """Adiciona coluna Reputação (delega à implementação compartilhada)."""
    from analysis import add_reputacao_column
    return add_reputacao_column(df)


def salvar_exportacao(df, export_cols, output_file, update_callback=None,
                      gerar_xlsx=True):
    """Sanitiza, grava o CSV e (quando viável) o Excel colorido.

    Tail comum aos pipelines de log de acesso e de interceptação, que antes
    tinham cópias separadas.

    O CSV é o artefato primário: se o Excel não puder ser gerado, o
    processamento continua e o motivo é reportado. Abortar uma execução de
    horas por causa de um formato de conveniência seria pior.

    Returns:
        (df_export, xlsx_path or None)
    """
    from helpers.atomic_io import atomic_write
    from validators import sanitize_dataframe_for_csv

    cols = [c for c in export_cols if c in df.columns]
    df_export = sanitize_dataframe_for_csv(df[cols])
    # Atômico: uma falha no meio deixava um CSV truncado com cara de
    # resultado completo — e é esse arquivo que o laudo cita.
    atomic_write(output_file,
                 lambda fh: df_export.to_csv(fh, index=False, sep=';'),
                 mode='w', encoding='utf-8-sig', newline='')

    xlsx_path = None
    if gerar_xlsx:
        candidato = os.path.splitext(output_file)[0] + '.xlsx'
        parcial = candidato + '.part'
        progresso = None
        if update_callback:
            def progresso(feitas, total):
                update_callback("Excel colorido: %s de %s linhas..."
                                % (_fmt_int(feitas), _fmt_int(total)))
        try:
            resumo = export_xlsx_colored(df_export, parcial, progress_callback=progresso)
            os.replace(parcial, candidato)
            xlsx_path = candidato
            if update_callback and resumo.get('abas', 1) > 1:
                # Continuação em abas precisa ser declarada: quem abre a
                # planilha vê a primeira aba cheia e não sabe que há mais.
                update_callback(f"Excel: {_fmt_int(resumo['linhas'])} linhas em "
                                f"{resumo['abas']} abas (limite de 1.048.575 por aba).")
        except ValueError as exc:
            logger.warning("Excel colorido não gerado: %s", exc)
            if update_callback:
                update_callback(f"Excel não gerado ({exc}) — CSV salvo normalmente.")
        except Exception as exc:
            logger.exception("Falha ao gerar Excel colorido")
            if update_callback:
                update_callback(f"Excel não gerado ({exc}) — CSV salvo normalmente.")
        finally:
            if os.path.exists(parcial):
                try:
                    os.remove(parcial)
                except OSError:
                    logger.exception("Falha ao remover o parcial %s", parcial)

    if update_callback:
        destino = f"{output_file} e .xlsx" if xlsx_path else output_file
        update_callback(f"Processamento finalizado. "
                        f"{len(df_export)} registros salvos em '{destino}'")
    return df_export, xlsx_path


# Mapa de cores por label de reputação. Em nível de módulo porque o teste de
# fidelidade visual e a legenda da interface se apoiam nesses valores.
XLSX_COLOR_MAP = {
    'Proxy / VPN / Tor': 'FECACA',   # vermelho claro
    'Datacenter / VPN':  'FEE2E2',   # vermelho mais claro
    'Cloud Pública':     'FEF3C7',   # amarelo claro
    'Hospedagem':        'FFEDD5',   # laranja claro
    'Rede Móvel':        'DBEAFE',   # azul claro
    'Residencial':       'DCFCE7',   # verde claro
}


def _larguras_colunas(df, columns):
    """Largura de cada coluna, estimada por uma amostra do topo."""
    larguras = []
    for col_name in columns:
        max_len = len(str(col_name))
        for val in df[col_name].head(100):
            if val is not None:
                max_len = max(max_len, len(str(val)))
        larguras.append(min(max_len + 2, 50))
    return larguras


def _nova_aba(wb, titulo, columns, larguras, estilo_cabecalho):
    """Cria uma aba com larguras e cabeçalho estilizado.

    No modo write_only não existe `ws.cell()`: o cabeçalho precisa ser uma
    lista de `WriteOnlyCell` já estilizada, escrita como primeira linha.
    """
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(titulo)
    for col_idx, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    font, fill, align = estilo_cabecalho
    cabecalho = []
    for col_name in columns:
        cell = WriteOnlyCell(ws, value=col_name)
        cell.font = font
        cell.fill = fill
        cell.alignment = align
        cabecalho.append(cell)
    ws.append(cabecalho)
    return ws


def export_xlsx_colored(df, output, *, sheet_row_limit=XLSX_SHEET_ROWS,
                        progress_callback=None):
    """Exporta o DataFrame para Excel com cores na coluna Reputação.

    Usa o modo `write_only` do openpyxl: cada linha é serializada para o XML no
    momento do `append` e as células são descartadas em seguida, então o
    consumo de memória não cresce com o número de linhas. O `Workbook` comum
    mantém um objeto `Cell` vivo por célula — 622 MB de heap só para 100 mil
    linhas x 15 colunas, medido — e era isso que impunha um teto de 100 mil.

    Acima de `sheet_row_limit` o resultado continua em abas adicionais
    (`Resultado (2)`, `Resultado (3)`, ...), porque o limite é do formato XLSX,
    não da implementação. Nada é omitido em silêncio: o retorno informa quantas
    linhas foram gravadas e em quantas abas, para o chamador declarar isso.

    O DataFrame é sanitizado aqui dentro contra fórmula injetada — o chamador
    não precisa (e não deve precisar) lembrar disso. O frame recebido não é
    modificado.

    Args:
        df: DataFrame a exportar.
        output: Caminho do arquivo ou buffer BytesIO.
        sheet_row_limit: Linhas de dados por aba. Só é reduzido em teste, para
            exercitar a divisão sem gerar um milhão de linhas.
        progress_callback: Chamado como `(feitas, total)` a cada 10 mil linhas.
            Este módulo não importa `streamlit`; renderizar o progresso é
            responsabilidade da página.

    Returns:
        dict: {'linhas': int, 'abas': int}

    Raises:
        ValueError: acima de `sheet_row_limit * XLSX_MAX_SHEETS` linhas.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import PatternFill, Font, Alignment

    from validators import sanitize_dataframe_for_csv

    total = 0 if df is None else len(df)
    teto = sheet_row_limit * XLSX_MAX_SHEETS
    if total > teto:
        # Antes de sanitizar: não vale copiar um frame de milhões de linhas
        # só para recusá-lo em seguida.
        raise ValueError(
            "Excel colorido suporta no máximo {} linhas em {} abas (recebidas "
            "{}). Filtre os dados ou use a exportação CSV.".format(
                _fmt_int(teto), XLSX_MAX_SHEETS, _fmt_int(total)))

    # Sanitização aqui dentro, e não no chamador: uma célula iniciada por
    # `=`, `+`, `-` ou `@` vira fórmula ao abrir a planilha, exatamente a
    # injeção que o projeto já bloqueia no CSV. Deixar isso a cargo de quem
    # chama é como o caminho da interface passou a exportar sem sanitizar
    # enquanto o do pipeline sanitizava. `sanitize_csv_value` prefixa com `'`,
    # que não está no conjunto perigoso — reaplicar sobre um frame já
    # sanitizado é inofensivo.
    df = sanitize_dataframe_for_csv(df)

    columns = list(df.columns)
    n_cols = len(columns)
    larguras = _larguras_colunas(df, columns)
    estilo_cabecalho = (
        Font(bold=True, color="FFFFFF"),
        PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid"),
        Alignment(horizontal="center"),
    )

    wb = Workbook(write_only=True)
    ws = _nova_aba(wb, "Resultado", columns, larguras, estilo_cabecalho)
    abas = 1

    # Um PatternFill por cor, não por linha: o laço antigo alocava um objeto
    # novo a cada linha colorida (centenas de milhares em datasets grandes).
    FILLS = {label: PatternFill(start_color=color, end_color=color, fill_type="solid")
             for label, color in XLSX_COLOR_MAP.items()}

    # Resolve a cor de cada linha de uma vez, vetorizado por valor distinto.
    row_fills = None
    if 'Reputação' in columns:
        def _fill_for(value):
            text = str(value)
            for label in XLSX_COLOR_MAP:
                if label in text:
                    return FILLS[label]
            return None
        distinct = {v: _fill_for(v) for v in df['Reputação'].unique()}
        row_fills = [distinct.get(v) for v in df['Reputação']]

    # Buffers de células reutilizáveis, um por cor (seis no máximo). O `append`
    # serializa a linha de forma síncrona, então dá para reaproveitar os mesmos
    # objetos — o que evita alocar milhões de WriteOnlyCell num caso real.
    buffers = {}

    def _buffer(aba, fill):
        chave = (id(aba), id(fill))
        buf = buffers.get(chave)
        if buf is None:
            buf = [WriteOnlyCell(aba) for _ in range(n_cols)]
            for cell in buf:
                cell.fill = fill
            buffers[chave] = buf
        return buf

    na_aba = 0
    for feitas, values in enumerate(df.itertuples(index=False, name=None), 1):
        if na_aba >= sheet_row_limit:
            abas += 1
            ws = _nova_aba(wb, "Resultado (%d)" % abas, columns, larguras,
                           estilo_cabecalho)
            na_aba = 0

        fill = row_fills[feitas - 1] if row_fills is not None else None
        if fill is None:
            # Caminho rápido do openpyxl: nenhum objeto Cell é criado.
            ws.append(values)
        else:
            buf = _buffer(ws, fill)
            for cell, val in zip(buf, values):
                cell.value = val
            ws.append(buf)
        na_aba += 1

        if progress_callback and feitas % 10_000 == 0:
            progress_callback(feitas, total)

    wb.save(output)
    if progress_callback:
        progress_callback(total, total)
    logger.info("Excel colorido: %d linhas em %d aba(s)", total, abas)
    return {'linhas': total, 'abas': abas}


def export_xlsx_to_disk(df, nome_base, progress_callback=None):
    """Gera o Excel colorido em disco, sob o sandbox `output/csv/`.

    Em datasets grandes o arquivo tem dezenas de MB. Mantê-lo como bytes em
    cache de sessão e empurrá-lo pelo websocket é o que derruba a conexão do
    Streamlit; servido de um handle de arquivo, o download é streaming.

    A gravação é atômica (`.part` + `os.replace`): uma falha no meio não deixa
    um `.xlsx` truncado com cara de export completo.

    Returns:
        (caminho, resumo) — `resumo` é o dict de `export_xlsx_colored`.
    """
    from validators import safe_output_path

    nome = nome_base if nome_base.lower().endswith('.xlsx') else nome_base + '.xlsx'
    destino = safe_output_path(nome, default_name='resultado_colorido.xlsx')
    parcial = destino + '.part'
    try:
        resumo = export_xlsx_colored(df, parcial, progress_callback=progress_callback)
        os.replace(parcial, destino)
    except Exception:
        try:
            if os.path.exists(parcial):
                os.remove(parcial)
        except OSError:
            logger.exception("Falha ao remover o parcial %s", parcial)
        raise
    return destino, resumo
